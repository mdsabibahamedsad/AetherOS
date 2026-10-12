"""Document parsers for supported file types."""

from __future__ import annotations

import csv
import json
import logging
import os
import io
import re
from pathlib import Path
from typing import Any, Callable, Optional

from aetheros_knowledge.models import SourceType

logger = logging.getLogger(__name__)


def parse_markdown(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def parse_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def parse_pdf(path: str) -> str:
    try:
        import PyPDF2
        text = []
        with open(path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text.append(page.extract_text() or "")
        return "\n".join(text)
    except ImportError:
        pass
    try:
        import pdfminer
        from pdfminer.high_level import extract_text
        return extract_text(path)
    except ImportError:
        pass
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return "\n".join(p.extract_text() or "" for p in pdf.pages)
    except ImportError:
        pass
    logger.warning("No PDF parser available, install PyPDF2 or pdfminer")
    return f"[PDF file: {os.path.basename(path)}]"


def parse_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    except ImportError:
        logger.warning("python-docx not installed")
        return f"[DOCX file: {os.path.basename(path)}]"


def parse_pptx(path: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(path)
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texts.append(shape.text)
        return "\n".join(texts)
    except ImportError:
        logger.warning("python-pptx not installed")
        return f"[PPTX file: {os.path.basename(path)}]"


def parse_csv(path: str) -> str:
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            return "\n".join(" | ".join(row) for row in reader)
    except Exception as e:
        logger.warning("CSV parse error: %s", e)
        return f"[CSV file: {os.path.basename(path)}]"


def parse_excel(path: str) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        texts = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            texts.append(f"--- Sheet: {sheet} ---")
            for row in ws.iter_rows(values_only=True):
                texts.append(" | ".join(str(c) for c in row if c is not None))
        return "\n".join(texts)
    except ImportError:
        pass
    try:
        import pandas as pd
        dfs = pd.read_excel(path, sheet_name=None)
        texts = []
        for name, df in dfs.items():
            texts.append(f"--- Sheet: {name} ---")
            texts.append(df.to_string())
        return "\n".join(texts)
    except ImportError:
        logger.warning("No Excel parser available")
        return f"[Excel file: {os.path.basename(path)}]"


def parse_json(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        data = json.load(f)
    return json.dumps(data, indent=2, default=str)


def parse_yaml(path: str) -> str:
    try:
        import yaml
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            data = yaml.safe_load(f)
        return yaml.dump(data, default_flow_style=False, allow_unicode=True)
    except ImportError:
        logger.warning("PyYAML not installed")
        return f"[YAML file: {os.path.basename(path)}]"


def parse_html(path: str) -> str:
    try:
        from bs4 import BeautifulSoup
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            soup = BeautifulSoup(f.read(), "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header"]):
            tag.decompose()
        return soup.get_text(separator="\n", strip=True)
    except ImportError:
        pass
    try:
        import html2text
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return html2text.html2text(f.read())
    except ImportError:
        logger.warning("No HTML parser available")
        return f"[HTML file: {os.path.basename(path)}]"


def parse_epub(path: str) -> str:
    try:
        import ebooklib
        from ebooklib import epub
        from bs4 import BeautifulSoup
        book = epub.read_epub(path)
        texts = []
        for item in book.get_items():
            if item.get_type() == ebooklib.ITEM_DOCUMENT:
                soup = BeautifulSoup(item.get_content(), "html.parser")
                texts.append(soup.get_text(separator="\n", strip=True))
        return "\n".join(texts)
    except ImportError:
        logger.warning("ebooklib not installed")
        return f"[EPUB file: {os.path.basename(path)}]"


PARSER_MAP: dict[SourceType, Callable] = {
    SourceType.MARKDOWN: parse_markdown,
    SourceType.TXT: parse_txt,
    SourceType.PDF: parse_pdf,
    SourceType.DOCX: parse_docx,
    SourceType.PPTX: parse_pptx,
    SourceType.CSV: parse_csv,
    SourceType.EXCEL: parse_excel,
    SourceType.JSON: parse_json,
    SourceType.YAML: parse_yaml,
    SourceType.HTML: parse_html,
    SourceType.EPUB: parse_epub,
}


def get_parser(source_type: SourceType) -> Callable:
    return PARSER_MAP.get(source_type, parse_txt)

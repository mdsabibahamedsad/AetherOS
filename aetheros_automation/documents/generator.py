from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import Any

from aetheros_automation.models.models import DocumentSpec

logger = logging.getLogger(__name__)

try:
    from docx import Document as DocxDocument
    from docx.shared import Inches, Pt

    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False
    logger.warning("python-docx not available; .docx generation disabled")

try:
    from fpdf import FPDF

    _HAS_FPDF = True
except ImportError:
    _HAS_FPDF = False
    logger.warning("fpdf2 not available; PDF generation disabled")

try:
    from openpyxl import Workbook

    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False
    logger.warning("openpyxl not available; Excel generation disabled")

try:
    from pptx import Presentation
    from pptx.util import Inches as PptxInches

    _HAS_PPTX = True
except ImportError:
    _HAS_PPTX = False
    logger.warning("python-pptx not available; PowerPoint generation disabled")


class DocumentGenerator:
    """Generate documents in multiple formats from structured specs."""

    @staticmethod
    async def generate(spec: DocumentSpec) -> str:
        """Dispatch to the correct generator based on spec.format."""
        fmt = spec.format.lower()
        output = spec.output_path or f"{spec.title}.{fmt}"

        generators = {
            "docx": lambda: DocumentGenerator.generate_word(spec.title, spec.content, output),
            "pdf": lambda: DocumentGenerator.generate_pdf(spec.title, spec.content, output),
            "xlsx": lambda: DocumentGenerator.generate_excel(spec.title, [], output),
            "pptx": lambda: DocumentGenerator.generate_powerpoint(spec.title, [], output),
            "md": lambda: DocumentGenerator.generate_markdown(spec.title, spec.content, output),
            "markdown": lambda: DocumentGenerator.generate_markdown(spec.title, spec.content, output),
            "html": lambda: DocumentGenerator.generate_html(spec.title, spec.content, output),
            "csv": lambda: DocumentGenerator.generate_csv([], [], output),
        }

        gen = generators.get(fmt)
        if gen:
            return await gen()

        if spec.template:
            return await DocumentGenerator.from_template(spec.template, {}, output)

        msg = f"Unsupported document format: {fmt}"
        raise ValueError(msg)

    @staticmethod
    async def generate_word(title: str, content: str, output_path: str) -> str:
        if not _HAS_DOCX:
            raise ImportError("python-docx is required to generate .docx files")
        out = Path(output_path).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        doc = DocxDocument()
        doc.add_heading(title, level=1)
        for paragraph in content.split("\n\n"):
            p = doc.add_paragraph(paragraph.strip())
            if paragraph.strip():
                p.style = doc.styles["Normal"]
        doc.save(str(out))
        logger.info("Generated Word document: %s", out)
        return str(out)

    @staticmethod
    async def generate_pdf(title: str, content: str, output_path: str) -> str:
        if not _HAS_FPDF:
            raise ImportError("fpdf2 is required to generate PDF files")
        out = Path(output_path).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 12)
        pdf.multi_cell(0, 8, content)
        pdf.output(str(out))
        logger.info("Generated PDF: %s", out)
        return str(out)

    @staticmethod
    async def generate_excel(title: str, data: list[list[Any]], output_path: str) -> str:
        if not _HAS_OPENPYXL:
            raise ImportError("openpyxl is required to generate Excel files")
        out = Path(output_path).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        for row_idx, row in enumerate(data, start=1):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        wb.save(str(out))
        logger.info("Generated Excel document: %s", out)
        return str(out)

    @staticmethod
    async def generate_powerpoint(title: str, slides: list[dict[str, Any]], output_path: str) -> str:
        if not _HAS_PPTX:
            raise ImportError("python-pptx is required to generate PowerPoint files")
        out = Path(output_path).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        prs = Presentation()
        slide_layout = prs.slide_layouts[1]
        for slide_def in slides:
            slide = prs.slides.add_slide(slide_layout)
            title_shape = slide.shapes.title
            content_shape = slide.placeholders[1]
            if title_shape:
                title_shape.text = slide_def.get("title", "")
            if content_shape:
                content_shape.text = slide_def.get("content", "")
        prs.save(str(out))
        logger.info("Generated PowerPoint: %s", out)
        return str(out)

    @staticmethod
    async def generate_markdown(title: str, content: str, output_path: str) -> str:
        out = Path(output_path).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        md = f"# {title}\n\n{content}"
        if not md.endswith("\n"):
            md += "\n"
        out.write_text(md, encoding="utf-8")
        logger.info("Generated Markdown: %s", out)
        return str(out)

    @staticmethod
    async def generate_html(title: str, content: str, output_path: str) -> str:
        out = Path(output_path).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; max-width: 800px; margin: 2em auto; padding: 0 1em; line-height: 1.6; }}
h1 {{ border-bottom: 2px solid #eee; padding-bottom: 0.3em; }}
</style>
</head>
<body>
<h1>{title}</h1>
{content}
</body>
</html>"""
        out.write_text(html, encoding="utf-8")
        logger.info("Generated HTML: %s", out)
        return str(out)

    @staticmethod
    async def generate_csv(data: list[list[Any]], headers: list[str] | None, output_path: str) -> str:
        out = Path(output_path).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        with out.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh)
            if headers:
                writer.writerow(headers)
            writer.writerows(data)
        logger.info("Generated CSV: %s", out)
        return str(out)

    @staticmethod
    async def from_template(template_path: str, variables: dict[str, Any], output_path: str) -> str:
        src = Path(template_path).resolve()
        raw = src.read_text(encoding="utf-8")
        rendered = raw
        for key, value in variables.items():
            rendered = rendered.replace(f"{{{{ {key} }}}}", str(value))
            rendered = rendered.replace(f"{{{{{key}}}}}", str(value))
        out = Path(output_path).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(rendered, encoding="utf-8")
        logger.info("Generated from template %s -> %s", src, out)
        return str(out)
